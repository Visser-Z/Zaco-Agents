"""Tests for the Nett writeback (into saved workbook rows) and period bounds
used by the delete-history action."""

import asyncio
from datetime import date

from openpyxl import Workbook

import app.main as main
from app import analytics, workbook
from app.columns import HEADERS, HEADER_ROW
from app.supabase_auth import User

USER = User(id="u-1", email="op@example.com", token="tok")


def _wb_with(rows):
    """A minimal workbook with headers and given (stm_no, nett) data rows."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for col, header in HEADERS.items():
        ws[f"{col}{HEADER_ROW}"] = header
    r = HEADER_ROW + 1
    for stm, nett in rows:
        ws[f"A{r}"] = 1000              # DN
        ws[f"E{r}"] = stm               # STM No
        ws[f"N{r}"] = nett              # Nett
        r += 1
    return wb


def test_write_netts_updates_matching_rows_by_stm_no():
    wb = _wb_with([(387517, None), (388486, None), (999999, 50.0)])
    updated = workbook.write_netts(wb, {387517: 2238.01, 388486: 7668.00})
    assert updated == 2
    rows = {r.stm_no: r.nett_total for r in workbook.read_rows(wb)}
    assert rows[387517] == 2238.01
    assert rows[388486] == 7668.00
    assert rows[999999] == 50.0          # untouched (no matching nett)


def test_write_netts_none_when_no_matches():
    wb = _wb_with([(111, None)])
    assert workbook.write_netts(wb, {222: 10.0}) == 0


def _wb_with_grosses(rows):
    """Rows carrying (stm_no, cartons sold, price), which is what the split needs."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for col, header in HEADERS.items():
        ws[f"{col}{HEADER_ROW}"] = header
    r = HEADER_ROW + 1
    for stm, sold, price in rows:
        ws[f"A{r}"], ws[f"E{r}"], ws[f"J{r}"], ws[f"L{r}"] = 1000, stm, sold, price
        r += 1
    return wb


def test_one_statement_over_several_products_splits_its_nett_by_gross():
    """An account sale settles every product on it with one figure. Writing that
    whole figure onto each row would multiply the money by the number of
    products -- here R3 000 would have become R6 000. Split by gross, which is
    exactly how the operator's own book does it."""
    wb = _wb_with_grosses([(375596, 100, 100.0), (375596, 200, 100.0)])
    assert workbook.write_netts(wb, {375596: 3000.0}) == 2
    netts = [r.nett_total for r in workbook.read_rows(wb)]
    assert netts == [1000.0, 2000.0]         # in the ratio of their gross
    assert sum(netts) == 3000.0


def test_the_split_adds_up_to_the_payment_exactly():
    """Three-way splits do not divide into cents, so the largest row absorbs the
    remainder. Otherwise the rows miss the payment by a cent and read as a
    discrepancy when nothing is wrong."""
    wb = _wb_with_grosses([(1, 1, 100.0), (1, 1, 100.0), (1, 1, 100.0)])
    workbook.write_netts(wb, {1: 100.0})
    netts = [r.nett_total for r in workbook.read_rows(wb)]
    assert sum(netts) == 100.0
    assert sorted(netts) == [33.33, 33.33, 33.34]


def test_a_statement_with_no_gross_to_split_by_shares_equally():
    """Nothing sold under it yet, so there is no ratio to use. An equal share is
    the only honest option, and it still adds up."""
    wb = _wb_with_grosses([(7, 0, 0.0), (7, 0, 0.0)])
    workbook.write_netts(wb, {7: 500.0})
    assert [r.nett_total for r in workbook.read_rows(wb)] == [250.0, 250.0]


# --- period bounds for delete --------------------------------------------

def test_period_bounds_month():
    assert analytics.period_bounds(month="2026-07") == ("2026-07-01", "2026-07-31")
    assert analytics.period_bounds(month="2026-02") == ("2026-02-01", "2026-02-28")


def test_period_bounds_week_and_precedence():
    lo, hi = analytics.period_bounds(week="2026-W31")
    assert date.fromisoformat(lo).isoweekday() == 1    # Monday
    assert date.fromisoformat(hi).isoweekday() == 7    # Sunday
    # week wins when both are given
    assert analytics.period_bounds(month="2026-07", week="2026-W31") == (lo, hi)


def test_period_bounds_none():
    assert analytics.period_bounds() == (None, None)


# --- delete endpoint guards ----------------------------------------------

def test_delete_requires_a_period(monkeypatch):
    called = False
    async def fake_delete(*a, **k):
        nonlocal called; called = True; return []
    monkeypatch.setattr(main, "db_delete", fake_delete)
    import pytest
    with pytest.raises(Exception):
        asyncio.run(main.delete_history(month=None, week=None, user=USER))
    assert not called


def test_delete_scopes_to_the_period(monkeypatch):
    captured = {}
    async def fake_delete(user, path, params):
        captured["path"] = path; captured["params"] = params
        return [{"id": 1}, {"id": 2}]
    monkeypatch.setattr(main, "db_delete", fake_delete)
    result = asyncio.run(main.delete_history(month="2026-07", week=None, user=USER))
    assert result == {"deleted": 2, "from": "2026-07-01", "to": "2026-07-31"}
    assert captured["path"] == "statements"
    assert captured["params"]["and"] == "(group_date.gte.2026-07-01,group_date.lte.2026-07-31)"
