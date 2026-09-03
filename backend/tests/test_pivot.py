"""The pivot table on its own sheet: native, refreshable, never stale.

openpyxl can model native pivots but does not document building one, so these
tests check the parts Excel actually validates. A pivot Excel rejects is worse
than no pivot: it reads the file as damaged and offers to repair it, which makes
the operator distrust the whole workbook.
"""

import zipfile
from datetime import date
from io import BytesIO

import openpyxl
import pytest

from app import pivot, workbook
from app.schemas import StatementRow


def _row(dn, stm, product="IMP Nect", agent="Farmers Trust", sold=10, nett=850.0):
    return StatementRow(source_file="june.csv", dn=dn, stm_no=stm, description=product,
                        market_agent=agent, qty_received=sold + 5, opening_stock=sold + 5,
                        cartons_sold=sold, price=100.0, nett_total=nett,
                        date=date(2026, 6, 1), invoice_date=date(2026, 6, 2), status="02.06")


def _built(rows=None):
    wb = workbook.new_workbook()
    workbook.append_rows(wb, rows or [
        _row(14776, 1), _row(14815, 2, "Strawberries"),
        _row(14828, 3, agent="Subtropico"), _row(14954, 4, "Imp Cherries 5kg", "Subtropico"),
    ], market_agent="Farmers Trust")
    assert workbook.refresh_pivot(wb) is True
    return wb


def _saved(wb):
    return zipfile.ZipFile(BytesIO(workbook.to_bytes(wb)))


def test_the_pivot_lands_on_its_own_sheet():
    wb = _built()
    assert pivot.SHEET_NAME in wb.sheetnames
    assert wb.sheetnames[0] == "Sheet1"          # the data sheet stays first


def test_the_file_carries_every_part_excel_requires():
    z = _saved(_built())
    names = z.namelist()
    assert "xl/pivotTables/pivotTable1.xml" in names
    assert "xl/pivotCache/pivotCacheDefinition1.xml" in names
    # And each one declared, or Excel treats the file as damaged.
    types = z.read("[Content_Types].xml").decode()
    assert "pivotTable" in types and "pivotCacheDefinition" in types


def test_the_relationship_chain_is_complete():
    """workbook -> cache, pivot sheet -> table, table -> cache. A break in any of
    the three is a repair prompt on open."""
    z = _saved(_built())
    wb_rels = z.read("xl/_rels/workbook.xml.rels").decode()
    assert "pivotCacheDefinition" in wb_rels
    assert "pivotCache" in z.read("xl/workbook.xml").decode()
    assert "pivotTable" in z.read("xl/worksheets/_rels/sheet2.xml.rels").decode()
    assert "pivotCacheDefinition" in z.read("xl/pivotTables/_rels/pivotTable1.xml.rels").decode()


def test_the_cache_is_empty_and_marked_stale():
    """The whole point. Shipping a cached copy of every figure means a second
    version of the truth that is wrong the moment a row is appended, so Excel is
    told to rebuild it from the sheet on open instead."""
    cache = _reopened(_built())[0].cache
    assert cache.refreshOnLoad is True
    assert cache.invalid is True
    assert cache.recordCount == 0
    assert "pivotCacheRecords" not in _saved(_built()).namelist()


def test_it_declares_a_version_excel_recognises():
    """openpyxl defaults these to 0, which is not a version: Excel reads the file
    as damaged and offers to repair it."""
    p = _reopened(_built())[0]
    assert p.createdVersion == pivot.CREATED_VERSION
    assert p.updatedVersion == pivot.CREATED_VERSION
    assert p.cache.refreshedVersion == pivot.CREATED_VERSION
    assert p.minRefreshableVersion == pivot.REFRESHABLE_FROM


def _reopened(wb):
    reloaded = openpyxl.load_workbook(BytesIO(workbook.to_bytes(wb)))
    return reloaded[pivot.SHEET_NAME]._pivots


def test_it_arranges_by_agent_and_product_and_measures_cartons_and_nett():
    p = _reopened(_built())[0]
    headers = [f.name for f in p.cache.cacheFields]
    assert [headers[f.x] for f in p.rowFields] == ["Market Agent", "Description"]
    assert [f.name for f in p.dataFields] == ["Sum of Cartons Sold", "Sum of Nett Total"]
    assert [headers[f.fld] for f in p.dataFields] == ["Cartons Sold", "Nett Total"]


def test_more_than_one_measure_puts_the_value_names_on_the_column_axis():
    """Field -2 is how that axis is spelled. Without it Excel shows one measure
    and silently drops the other."""
    p = _reopened(_built())[0]
    assert -2 in [f.x for f in p.colFields]


def test_the_source_range_covers_every_row_and_grows_on_the_next_append():
    wb = _built([_row(14776, 1), _row(14815, 2)])
    assert _reopened(wb)[0].cache.cacheSource.worksheetSource.ref == "A1:U3"

    workbook.append_rows(wb, [_row(14954, 3)], market_agent="Farmers Trust")
    workbook.refresh_pivot(wb)
    p = _reopened(wb)[0]
    assert p.cache.cacheSource.worksheetSource.ref == "A1:U4"
    assert p.cache.cacheSource.worksheetSource.sheet == "Sheet1"


def test_rebuilding_leaves_exactly_one_pivot_sheet():
    """Appending twice must not stack up Pivot, Pivot1, Pivot2."""
    wb = _built()
    for _ in range(3):
        workbook.refresh_pivot(wb)
    assert [s for s in wb.sheetnames if s.startswith(pivot.SHEET_NAME)] == [pivot.SHEET_NAME]
    assert len(_reopened(wb)) == 1


def test_no_pivot_over_an_empty_sheet():
    """A pivot with no rows under it is a repair prompt, not a convenience."""
    wb = workbook.new_workbook()
    assert workbook.refresh_pivot(wb) is False
    assert pivot.SHEET_NAME not in wb.sheetnames


def test_a_sheet_without_the_fields_gets_no_pivot_rather_than_a_broken_one():
    """Nothing to arrange by, or nothing to measure, so there is no pivot to
    build. Writing one anyway would point at fields that do not exist."""
    wb = openpyxl.Workbook()          # bare, not the template
    ws = wb.active
    ws["A1"], ws["B1"] = "DN", "Something else"
    ws["A2"], ws["B2"] = 14776, 1
    assert pivot.add_pivot(wb, ws, 1, 2) is False
    assert pivot.SHEET_NAME not in wb.sheetnames


def test_a_sheet_with_rows_but_no_measure_gets_no_pivot():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["B1"] = "Market Agent", "Description"
    ws["A2"], ws["B2"] = "Farmers Trust", "IMP Nect"
    assert pivot.add_pivot(wb, ws, 1, 2) is False


def test_a_pivot_failure_never_costs_the_save(monkeypatch):
    monkeypatch.setattr(pivot, "add_pivot", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("boom")))
    wb = workbook.new_workbook()
    workbook.append_rows(wb, [_row(14776, 1)], market_agent="Farmers Trust")
    assert workbook.refresh_pivot(wb) is False
    assert workbook.to_bytes(wb)                  # still saves


def test_the_data_sheet_is_untouched_by_the_pivot():
    """The pivot reads the sheet; it must never write to it."""
    wb = _built()
    ws = wb["Sheet1"]
    before = [[c.value for c in r] for r in ws.iter_rows()]
    workbook.refresh_pivot(wb)
    assert [[c.value for c in r] for r in ws.iter_rows()] == before
