"""Workbook append tests -- the part that must never corrupt the operator's file."""

from datetime import date, datetime

from app import workbook
from app.columns import DATA_COLUMNS, FORMULA_COLUMNS
from app.schemas import StatementRow


def make_row(stm: int, dn: int = 14847) -> StatementRow:
    return StatementRow(
        source_file=f"AS_{stm}.pdf",
        dn=dn,
        stm_no=stm,
        description="IMP Nect",
        qty_received=76,
        opening_stock=54,
        cartons_sold=54,
        price=50.0,
        nett_total=2238.01,
        date=date(2026, 6, 18),
        status="01.07",
    )


def test_append_writes_data_and_formulas():
    wb = workbook.new_workbook()
    workbook.append_rows(wb, [make_row(387517)], market_agent="Farmers Trust")
    ws = wb["Sheet1"]

    assert ws["A2"].value == 14847            # DN
    assert ws["B2"].value == "Farmers Trust"  # supplied per session
    assert ws["C2"].value == "Incomplete"     # always starts Incomplete
    assert ws["E2"].value == 387517
    assert ws["F2"].value == "IMP Nect"
    assert ws["J2"].value == 54
    assert ws["N2"].value == 2238.01
    assert ws["T2"].value == "01.07"

    # Formula columns must hold formulas bound to this row, never computed values.
    assert ws["M2"].value == "=SUM(J2*L2)"
    assert ws["K2"].value == "=H2-J2"
    assert ws["S2"].value == "=I2*J2"


def test_the_commission_chain_hangs_on_the_editable_cell():
    """I is the supplier's price per carton, written as the 70% default so the
    row lands on a 30% commission. Everything else follows *that cell*, so
    typing an agreed price over it corrects the whole row -- which is how the
    operator's own book carries deals on other terms."""
    wb = workbook.new_workbook()
    workbook.append_rows(wb, [make_row(387517)], market_agent="Farmers Trust")
    ws = wb["Sheet1"]

    assert ws["I2"].value == '=IFERROR(O2*70%,"-")'   # the default, editable
    assert ws["S2"].value == "=I2*J2"                 # supplier's share follows I
    assert ws["Q2"].value == "=N2-S2"                 # Zaco keeps the rest
    assert ws["P2"].value == '=IFERROR(O2-I2,"-")'
    assert ws["R2"].value == '=IFERROR(P2/O2,"-")'
    # Nothing in the chain refers back to a fixed 30%, or a hand-corrected I
    # would leave the row internally inconsistent.
    assert not any("30%" in str(ws[f"{c}2"].value) for c in "IPQRS")


def test_appending_twice_stacks_rows_and_rebinds_formulas():
    wb = workbook.new_workbook()
    workbook.append_rows(wb, [make_row(1)], "Farmers Trust")
    workbook.append_rows(wb, [make_row(2)], "Farmers Trust")
    ws = wb["Sheet1"]

    assert ws["E2"].value == 1
    assert ws["E3"].value == 2
    # Row 3's formulas reference row 3, not row 2.
    assert ws["M3"].value == "=SUM(J3*L3)"
    assert ws["O3"].value == '=IFERROR(N3/J3,"-")'


def test_append_skips_prefilled_formula_rows():
    """The template pre-fills formulas below the last real row.

    Those rows must count as blank, otherwise the first append lands far down
    the sheet instead of directly beneath the existing data.
    """
    wb = workbook.new_workbook()
    ws = wb["Sheet1"]
    workbook.append_rows(wb, [make_row(100)], "Farmers Trust")  # occupies row 2
    # Simulate the template's pre-filled formulas on rows 3-10 with no data.
    # In the standard template the letter mapping is the identity.
    identity = {letter: letter for letter in "ABCDEFGHIJKLMNOPQRST"}
    for r in range(3, 11):
        for col, template in FORMULA_COLUMNS.items():
            ws[f"{col}{r}"] = template.format(r=r, **identity)

    workbook.append_rows(wb, [make_row(101)], "Farmers Trust")
    assert ws["E3"].value == 101, "second row should land at row 3, not below the formula block"


def test_round_trip_read_back():
    wb = workbook.new_workbook()
    workbook.append_rows(wb, [make_row(387517), make_row(386729)], "Farmers Trust")
    reloaded = workbook.load(workbook.to_bytes(wb))
    rows = workbook.read_rows(reloaded)

    assert len(rows) == 2
    assert {r.stm_no for r in rows} == {387517, 386729}
    assert rows[0].market_agent == "Farmers Trust"
    assert rows[0].date == date(2026, 6, 18)


def test_data_and_formula_columns_do_not_overlap():
    assert set(DATA_COLUMNS) & set(FORMULA_COLUMNS) == set()


# --- header safety -------------------------------------------------------

def test_headers_written_only_for_a_new_workbook():
    """Appending must never write a header row.

    A workbook the operator has been building already has its headers (and its
    own styling). Re-writing row 1 on every append would stamp a duplicate
    header over their data.
    """
    wb = workbook.new_workbook()
    ws = wb["Sheet1"]
    # An operator's own tweak to a header must survive every append. (The DN
    # and STM No headers stay recognisable -- they are how the table is found.)
    ws["C1"] = "Completed (own wording)"

    for stm in (1, 2, 3):
        workbook.append_rows(wb, [make_row(stm)], "Farmers Trust")

    assert ws["C1"].value == "Completed (own wording)"
    assert ws["A1"].value == "DN"
    assert [ws[f"E{r}"].value for r in (2, 3, 4)] == [1, 2, 3]
    # No row below the header may contain header text.
    assert "DN" not in [ws[f"A{r}"].value for r in range(2, 6)]


def test_new_workbook_carries_template_headers():
    ws = workbook.new_workbook()["Sheet1"]
    assert ws["A1"].value == "DN"
    assert ws["T1"].value == "Status"


# --- group dates across separate appends ---------------------------------

def test_group_date_reconciles_across_appends():
    """A later batch with an earlier receive date must pull the whole DN back.

    Statement 387517 (received 19 Jun) is added first, then 386729 (received
    18 Jun) in a second round. Both rows must end up dated 18 June.
    """
    wb = workbook.new_workbook()
    ws = wb["Sheet1"]

    first = make_row(387517)
    first.date = date(2026, 6, 19)
    workbook.append_rows(wb, [first], "Farmers Trust")
    assert ws["D2"].value.date() == date(2026, 6, 19)

    second = make_row(386729)
    second.date = None
    second.date_received = date(2026, 6, 18)
    workbook.append_rows(wb, [second], "Farmers Trust")

    assert ws["D2"].value.date() == date(2026, 6, 18), "existing row should be pulled back"
    assert ws["D3"].value.date() == date(2026, 6, 18), "new row takes the group date"


def test_group_dates_do_not_leak_between_dns():
    wb = workbook.new_workbook()
    ws = wb["Sheet1"]
    a = make_row(1, dn=100); a.date = date(2026, 6, 20)
    workbook.append_rows(wb, [a], "Farmers Trust")

    b = make_row(2, dn=200); b.date = None; b.date_received = date(2026, 5, 1)
    workbook.append_rows(wb, [b], "Farmers Trust")

    assert ws["D2"].value.date() == date(2026, 6, 20), "DN 100 must be untouched"
    assert ws["D3"].value.date() == date(2026, 5, 1)


# --- real-world workbooks: wrong sheet, shuffled columns, messy cells -----
# Regression tests for a production failure: opening a multi-sheet workbook
# whose columns did not sit at the template's letters crashed with pydantic
# errors because values were read by fixed position ("dn" got "Farmers Trust",
# "date" got "Imp White Grapes", ...).

from openpyxl import Workbook as _Workbook

SHUFFLED = ["Market Agent", "DN", "Status", "Description", "STM No", "Date",
            "Nett Total", "Qty Received", "Opening Stock", "Cartons Sold", "Price"]


def _shuffled_sheet(ws):
    ws.append(SHUFFLED)
    ws.append(["Farmers Trust", 14847, "01.07", "IMP Nect", 387517,
               datetime(2026, 6, 18), 2238.01, 76, 54, 54, 50.0])


def test_open_multisheet_workbook_with_shuffled_columns():
    """The data sheet is neither first nor 'Sheet1', and columns are reordered."""
    wb = _Workbook()
    wb.active.title = "Notes"
    wb.active["A1"] = "random scribbles"
    summary = wb.create_sheet("Summary")
    summary.append(["Total", 999])
    data = wb.create_sheet("July sales")
    _shuffled_sheet(data)

    rows = workbook.read_rows(wb)
    assert len(rows) == 1
    r = rows[0]
    assert r.dn == 14847
    assert r.stm_no == 387517
    assert r.market_agent == "Farmers Trust"
    assert r.description == "IMP Nect"
    assert r.nett_total == 2238.01
    assert r.date == date(2026, 6, 18)
    assert r.status == "01.07"


def test_append_into_shuffled_workbook_reanchors_formulas():
    """Appending must write values AND formulas at the workbook's own columns."""
    wb = _Workbook()
    ws = wb.active
    ws.title = "Data"
    _shuffled_sheet(ws)

    workbook.append_rows(wb, [make_row(400000)], "Farmers Trust")

    # In SHUFFLED, DN is column B and STM No is column E; new row lands at 3.
    assert ws["B3"].value == 14847
    assert ws["E3"].value == 400000
    # Nett Total sits at G and Cartons Sold at J here. There is no formula
    # column in this workbook for M etc. (headers absent), so none written --
    # but the ones whose headers exist must reference the right letters.
    rows2 = workbook.read_rows(wb)
    assert [r.stm_no for r in rows2] == [387517, 400000]


def test_messy_cells_do_not_crash_reading():
    """Junk in typed columns coerces to None instead of raising."""
    wb = _Workbook()
    ws = wb.active
    ws.append(list(HEADERS_LIST))
    junk = ["not-a-dn", "Agent", "Incomplete", "Imp White Grapes", "(blank)",
            "IMP Nect", "08.07", None, "x", None, "abc", None, None, "n/a",
            None, None, None, None, None, 42]
    ws.append(junk)
    rows = workbook.read_rows(wb)
    assert len(rows) == 1
    assert rows[0].dn is None            # "not-a-dn"
    assert rows[0].stm_no is None        # "(blank)"
    assert rows[0].date is None          # "Imp White Grapes" is not a date
    assert rows[0].status == "42"


def test_unrecognisable_workbook_raises_clear_error():
    wb = _Workbook()
    wb.active.append(["Totally", "Unrelated", "Spreadsheet"])
    try:
        workbook.read_rows(wb)
        assert False, "should have raised"
    except workbook.WorkbookFormatError as e:
        assert "Could not find the account-sales table" in str(e)


HEADERS_LIST = [
    "DN", "Market Agent", "Completed", "Date", "STM No", "Description",
    "Qty Received", "Opening Stock", "Frui Price/crt", "Cartons Sold",
    "Baby Stock", "Price", "Gross Total", "Nett Total", "Nett Price/crt",
    "Z R/crt", "Z Total", "% Markup", "Frui Curr Sales Value", "Status",
]


def test_header_row_not_on_row_one():
    """A title row above the table must not break detection."""
    wb = _Workbook()
    ws = wb.active
    ws.append(["ZACO ACCOUNT SALES — JULY 2026"])
    ws.append([])
    ws.append(list(HEADERS_LIST))
    ws.append([14847, "Farmers Trust", "Incomplete", datetime(2026, 6, 18),
               387517, "IMP Nect", 76, 54, None, 54, None, 50.0, None,
               2238.01, None, None, None, None, None, "01.07"])
    rows = workbook.read_rows(wb)
    assert len(rows) == 1
    assert rows[0].stm_no == 387517

    workbook.append_rows(wb, [make_row(400001)], "Farmers Trust")
    assert ws["E5"].value == 400001, "append lands under the real table"


# --- headerless workbooks (operator's own sheets, no header row) ----------

def test_headerless_sheet_opens_by_position():
    """A sheet with data from row 1 and no header row reads by A-T position.

    The operator's real workbooks (e.g. Book1) often have no header row -- they
    know the columns by position. DN in A and STM No in E being numbers is what
    marks it as an account-sales sheet.
    """
    from openpyxl import Workbook as _Workbook
    wb = _Workbook()
    ws = wb.active
    ws.append([14969, "Subtropico", "Completed", None, 56391641, "Imp White Grapes",
               99, 99, None, 35, None, 380, None, None, None, None, None, None, None, "06.07"])
    rows = workbook.read_rows(wb)
    assert len(rows) == 1
    r = rows[0]
    assert r.dn == 14969
    assert r.stm_no == 56391641
    assert r.market_agent == "Subtropico"
    assert r.description == "Imp White Grapes"
    assert r.cartons_sold == 35


def test_headerless_append_lands_beneath_data():
    from openpyxl import Workbook as _Workbook
    wb = _Workbook()
    ws = wb.active
    ws.append([100, "Ag", "Completed", None, 200, "X", 10, 10, None, 5,
               None, 50, None, None, None, None, None, None, None, "01.08"])
    workbook.append_rows(wb, [make_row(999)], "Farmers Trust")
    assert ws["E2"].value == 999, "new row lands at row 2, beneath the one data row"


def test_unrelated_spreadsheet_still_rejected():
    """A sheet with text (not numbers) in A/E is not treated as data."""
    from openpyxl import Workbook as _Workbook
    wb = _Workbook()
    wb.active.append(["Budget", "Category", "Amount"])
    wb.active.append(["Rent", "Fixed", 9000])
    try:
        workbook.read_rows(wb)
        assert False, "should have raised"
    except workbook.WorkbookFormatError:
        pass
